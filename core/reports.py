"""Infrastructure commune d'export (Excel/PDF) et d'entete de rapport pour les listes/tableaux.

Chaque liste (transactions MVOLA, transactions PAMF, ecarts, lignes d'un rapprochement) definit
ses propres colonnes (`[(libelle, fonction_extraction), ...]`) et construit un `meta` (titre,
module source, filtres actifs, nombre de lignes, auteur/horodatage) via `build_meta`. L'export
reutilise exactement le meme queryset filtre que l'ecran (pas de pagination), cf. demande
utilisateur : le rapport exporte doit correspondre a ce qui est affiche/filtre.
"""

from datetime import datetime
from decimal import Decimal

import openpyxl
from django.http import Http404, HttpResponse
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa


def build_meta(request, *, titre, module_source, filtres, nb_lignes, export_excel_url=None, export_pdf_url=None):
    return {
        'titre': titre,
        'module_source': module_source,
        'filtres': filtres,
        'nb_lignes': nb_lignes,
        'genere_par': str(request.user),
        'genere_le': timezone.now(),
        'export_excel_url': export_excel_url,
        'export_pdf_url': export_pdf_url,
    }


def decrire_filtres(form):
    """Resume textuel des filtres actifs d'un formulaire de filtre (label -> valeur affichee)."""
    filtres = []
    if not form.is_bound or not form.is_valid():
        return filtres
    for name, valeur in form.cleaned_data.items():
        if valeur in (None, ''):
            continue
        field = form.fields[name]
        if hasattr(field, 'choices'):
            valeur = dict(field.choices).get(valeur, valeur)
        filtres.append((field.label, valeur))
    return filtres


def urls_export(request, url_name, *url_args):
    """Construit les URLs d'export excel/pdf en reprenant la querystring de filtres actuelle."""
    qs = request.GET.urlencode()
    urls = {}
    for format_export in ('excel', 'pdf'):
        url = reverse(url_name, args=[*url_args, format_export])
        if qs:
            url = f'{url}?{qs}'
        urls[format_export] = url
    return urls['excel'], urls['pdf']


def _valeur_excel(valeur):
    if valeur is None:
        return ''
    if isinstance(valeur, Decimal):
        return float(valeur)
    if isinstance(valeur, datetime):
        if timezone.is_aware(valeur):
            valeur = timezone.localtime(valeur)
        return valeur.replace(tzinfo=None)
    return valeur


def export_excel(meta, colonnes, objets, nom_fichier):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Rapport'

    ws.append([meta['titre']])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f"Module source : {meta['module_source']}"])
    if meta['filtres']:
        for label, valeur in meta['filtres']:
            ws.append([f'{label} : {valeur}'])
    else:
        ws.append(['Filtres appliques : aucun'])
    ws.append([f"Nombre de lignes : {meta['nb_lignes']}"])
    ws.append([f"Genere par {meta['genere_par']} le {timezone.localtime(meta['genere_le']).strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])

    entete_row = ws.max_row + 1
    ws.append([label for label, _ in colonnes])
    for cell in ws[entete_row]:
        cell.font = Font(bold=True)

    for objet in objets:
        ws.append([_valeur_excel(getter(objet)) for _, getter in colonnes])

    for idx in range(1, len(colonnes) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 22

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}.xlsx"'
    wb.save(response)
    return response


def export_pdf(meta, colonnes, objets, nom_fichier):
    lignes = [[getter(objet) for _, getter in colonnes] for objet in objets]
    html = render_to_string('core/rapport_pdf.html', {
        'meta': meta,
        'colonnes_labels': [label for label, _ in colonnes],
        'lignes': lignes,
    })
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}.pdf"'
    pisa_status = pisa.CreatePDF(html, dest=response)
    if pisa_status.err:
        return HttpResponse('Erreur lors de la generation du PDF.', status=500)
    return response


def exporter(format_export, meta, colonnes, objets, nom_fichier):
    if format_export == 'excel':
        return export_excel(meta, colonnes, objets, nom_fichier)
    if format_export == 'pdf':
        return export_pdf(meta, colonnes, objets, nom_fichier)
    raise Http404("Format d'export inconnu.")
